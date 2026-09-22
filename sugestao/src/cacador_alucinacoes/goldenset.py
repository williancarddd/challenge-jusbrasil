"""Leitura do gabarito e dos documentos de entrada.

Duas coisas do `.xlsx` que mordem:

1. Toda coluna numérica volta como float — `nivel` é `1.0`, `inicio` é `469.0`,
   `id_canonico` é `5665364632.0`. Nada de especial no `id_canonico`: é uniforme, e
   como os ids cabem em 2^53 o float64 é exato. Basta `int(float(...))`.
2. As 129 linhas sem `id_canonico` voltam com **8 células em vez de 9** — a última é
   truncada. Um `zip(header, row)` descartaria a chave em silêncio, então as linhas são
   preenchidas até a largura do cabeçalho antes de virar dicionário.

Os nomes de coluna do gabarito e do contrato de submissão são em português
(`inicio`, `fim`, `trecho`, `tipo`, `classificacao`, `id_canonico`). Eles ficam
confinados à fronteira de leitura e serialização; dentro do código os campos são
`start`, `end`, `excerpt`, `kind`, `label`, `canonical_id`.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
INPUT_DIR = ROOT / "input"
GOLDENSET_PATH = ROOT / "data" / "goldenset.xlsx"


@dataclass(frozen=True)
class Citation:
    level: int
    document_id: str
    citation_id: str
    start: int
    end: int
    excerpt: str
    kind: str
    label: str
    canonical_id: str | None

    @property
    def weight(self) -> int:
        """Nível 1 pesa 1x e nível 2 pesa 2x na nota final."""
        return self.level


def _as_int(value) -> int:
    return int(float(value))


def load_goldenset(path: Path = GOLDENSET_PATH) -> list[Citation]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = [str(column) for column in next(rows)]

    citations = []
    for row in rows:
        padded = tuple(row) + (None,) * (len(header) - len(row))
        record = dict(zip(header, padded))
        if record["documento_id"] is None:
            continue
        canonical = record["id_canonico"]
        citations.append(
            Citation(
                level=_as_int(record["nivel"]),
                document_id=str(record["documento_id"]),
                citation_id=str(record["citacao_id"]),
                start=_as_int(record["inicio"]),
                end=_as_int(record["fim"]),
                excerpt=str(record["trecho"]),
                kind=str(record["tipo"]),
                label=str(record["classificacao"]),
                # Desde o e-mail de 28/08 o campo é um doc_id único, não mais conjunto.
                canonical_id=None if canonical in (None, "") else str(_as_int(canonical)),
            )
        )
    return citations


def load_documents(directory: Path = INPUT_DIR) -> dict[str, str]:
    return {
        path.stem: path.read_text(encoding="utf-8")
        for path in sorted(directory.glob("*.txt"))
    }


def literal_excerpt(document: str, start: int, end: int) -> str:
    """O `trecho` da submissão sai daqui — do arquivo cru, nunca da saída da LLM.

    O gabarito define `trecho` como cópia literal de `texto[inicio:fim]`, e 59 dos 225
    trechos contêm quebra de linha. No texto normalizado essa quebra virou espaço, então
    copiar da resposta da LLM produziria um `trecho` diferente do esperado.
    """
    return document[start:end]
